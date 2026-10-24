"""将指定版本的完整对话导出为 Word，不依赖网页折叠状态。"""

from io import BytesIO

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from podcast_workbench.models import Revision


def export_script(revision: Revision) -> bytes:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.85)
    for name in ("Normal", "Title", "Heading 1"):
        style = document.styles[name]
        style.font.name = "Microsoft YaHei"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
    normal = document.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.4
    normal.paragraph_format.space_after = Pt(9)
    document.add_heading(revision.script.title, 0)
    document.add_paragraph(f"对话脚本 · 第 {revision.number} 版")
    document.add_paragraph(
        f"主持人：{revision.settings.host_name}    嘉宾：{revision.settings.guest_name}"
    )
    for turn in revision.script.turns:
        speaker = (revision.settings.host_name if turn.role == "host"
                   else revision.settings.guest_name)
        # 称呼和正文同段，长回答可自然跨页，避免整段被挤到下一页。
        paragraph = document.add_paragraph()
        paragraph.add_run(speaker + "：").bold = True
        paragraph.add_run(turn.text)
    document.core_properties.title = revision.script.title
    document.core_properties.author = ""
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()
