from io import BytesIO
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient
from test_service import make_service
from test_store import example_script

from podcast_workbench.api import create_app
from podcast_workbench.models import ProgramSettings


def test_home_serves_editorial_ui_without_demo_results(tmp_path: Path) -> None:
    service, _, _ = make_service(tmp_path)
    with TestClient(create_app(service.config, service)) as client:
        home = client.get("/")
        assert home.status_code == 200
        assert "从一篇文章，到两个人的声音" in home.text
        assert "data-state=" not in home.text
        assert "设计示例" not in home.text
        assert "PRIVATE" not in home.text
        assert client.get("/static/app.js").status_code == 200
        assert client.get("/static/styles.css").status_code == 200
        assert client.get("/health").json()["frontend_connected"] is True


def test_static_mount_does_not_expose_project_files(tmp_path: Path) -> None:
    service, _, _ = make_service(tmp_path)
    with TestClient(create_app(service.config, service)) as client:
        paths = ("/.env", "/config.yaml", "/static/../config.yaml", "/static/%2e%2e/config.yaml")
        for path in paths:
            result = client.get(path)
            assert result.status_code == 404


def test_document_export_uses_complete_revision_and_custom_names(tmp_path: Path) -> None:
    service, _, _ = make_service(tmp_path)
    settings = ProgramSettings(host_name="小林", guest_name="阿周")
    session_id = service.store.create_session("原材料", settings)
    revision = service.store.save_revision(session_id, example_script(), settings=settings)
    with TestClient(create_app(service.config, service)) as client:
        result = client.get(f"/revisions/{revision.id}/document")
        assert result.status_code == 200
        assert "wordprocessingml.document" in result.headers["content-type"]
        assert ".docx" in result.headers["content-disposition"]
        document = Document(BytesIO(result.content))
        text = "\n".join(p.text for p in document.paragraphs)
        assert revision.script.title in text
        assert "小林" in text and "阿周" in text
        for turn in revision.script.turns:
            assert turn.text in text
        # 后续新版本不能使旧版导出悄悄变成当前版本。
        service.store.save_revision(session_id, example_script(), settings=ProgramSettings())
        old = client.get(f"/revisions/{revision.id}/document")
        assert "小林" in "\n".join(p.text for p in Document(BytesIO(old.content)).paragraphs)
        missing = client.get("/revisions/11111111-1111-4111-8111-111111111111/document")
        assert missing.status_code == 404
